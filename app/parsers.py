"""Extract plain text from uploaded resumes (PDF, DOCX, TXT)."""

from __future__ import annotations

import io


class UnsupportedFileError(ValueError):
    pass


def extract_text(filename: str, data: bytes) -> str:
    name = (filename or "").lower()
    if name.endswith(".pdf") or data[:4] == b"%PDF":
        return _pdf(data)
    if name.endswith(".docx") or data[:2] == b"PK":
        return _docx(data)
    if name.endswith((".txt", ".md")) or not name:
        try:
            return data.decode("utf-8")
        except UnicodeDecodeError:
            return data.decode("latin-1")
    raise UnsupportedFileError("Upload a PDF, DOCX or TXT resume.")


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
    except Exception as exc:  # pypdf raises several error types for broken files
        raise UnsupportedFileError("Could not read this PDF.") from exc
    if not text.strip():
        raise UnsupportedFileError("This PDF has no extractable text (is it a scanned image?).")
    return text


def _docx(data: bytes) -> str:
    from docx import Document

    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        raise UnsupportedFileError("Could not read this DOCX file.") from exc
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:  # many resume templates put content in tables
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)
